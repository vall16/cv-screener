"""API FastAPI del progetto cv-screener."""
import io
import os
import shutil
import subprocess
import zipfile
from datetime import datetime
from pathlib import Path

import markdown
from docx import Document as DocxDocument
from fastapi import FastAPI, File, HTTPException, UploadFile
from pydantic import BaseModel
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from xhtml2pdf import pisa

import screener

BASE_DIR = Path(__file__).resolve().parent.parent
FRONTEND_DIST = BASE_DIR / "frontend" / "dist" / "cv-screener"
if (FRONTEND_DIST / "browser").is_dir():
    FRONTEND_DIST = FRONTEND_DIST / "browser"

app = FastAPI(title="cv-screener", version="1.0.0")

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

_ALLOWED_EXT = {".pdf", ".txt", ".docx"}
_MAX_UPLOAD_MB = 15


def _docx_to_text(data: bytes) -> str:
    doc = DocxDocument(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs if p.text.strip()]
    for table in doc.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text.strip() for cell in row.cells))
    return "\n".join(parts)


@app.get("/api/health")
def health():
    return {"status": "ok"}


# ── CV ────────────────────────────────────────────────
@app.get("/api/cvs")
def api_list_cvs():
    return {"cvs": screener.list_cvs()}


@app.post("/api/cvs/upload")
async def api_upload_cvs(files: list[UploadFile] = File(...)):
    saved = []
    errors = []
    for f in files:
        name = Path(f.filename or "").name
        if not name or Path(name).suffix.lower() not in _ALLOWED_EXT:
            errors.append((name or "file") + ": formato non supportato (solo PDF/TXT/DOCX)")
            continue
        if f.size and f.size > _MAX_UPLOAD_MB * 1024 * 1024:
            errors.append(f"{name}: oltre {_MAX_UPLOAD_MB} MB")
            continue
        chunk = await f.read()
        if len(chunk) > _MAX_UPLOAD_MB * 1024 * 1024:
            errors.append(f"{name}: oltre {_MAX_UPLOAD_MB} MB")
            continue
        screener.DEFAULT_CV_DIR.mkdir(parents=True, exist_ok=True)
        if Path(name).suffix.lower() == ".docx":
            txt_name = Path(name).stem + ".txt"
            try:
                text = _docx_to_text(chunk)
                (screener.DEFAULT_CV_DIR / txt_name).write_text(text, encoding="utf-8")
                saved.append(txt_name)
            except Exception:
                errors.append(f"{name}: conversione DOCX fallita")
        else:
            (screener.DEFAULT_CV_DIR / name).write_bytes(chunk)
            saved.append(name)
    return {"saved": saved, "errors": errors, "cvs": screener.list_cvs()}


@app.delete("/api/cvs/all")
def api_delete_all_cvs():
    deleted = []
    for cv in screener.list_cvs():
        if screener.delete_cv(cv):
            deleted.append(cv)
    return {"deleted": deleted, "count": len(deleted)}


@app.delete("/api/cvs/{name}")
def api_delete_cv(name: str):
    if not screener.delete_cv(name):
        raise HTTPException(status_code=404, detail="CV non trovato")
    return {"ok": True}


# ── Jobs ──────────────────────────────────────────────
def _zip_cvs_stream(names: list[str], chunk_size: int = 256 * 1024):
    """Genera lo ZIP dei CV in memoria e lo spacca in blocchi.

    Niente file temporaneo: sul server non viene scritto nulla, l'archivio vive
    solo nella memoria di questo processo e viene consumato via via dal client,
    che lo scrive sul proprio disco. compresslevel=1 perche' i PDF sono gia'
    compressi e deflare di piu' non ci guadagnerebbe.
    """
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED, compresslevel=1) as zf:
        for name in names:
            src = screener.DEFAULT_CV_DIR / name
            if src.is_file():
                zf.write(src, arcname=name)
    buf.seek(0)
    while True:
        chunk = buf.read(chunk_size)
        if not chunk:
            break
        yield chunk


@app.get("/api/cvs/export")
def api_export_cvs():
    """Scarica tutti i CV in uno ZIP. Pensato per il caso server: il client li porta in locale."""
    names = screener.list_cvs()
    if not names:
        raise HTTPException(status_code=404, detail="Nessun CV da esportare")
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    return StreamingResponse(
        _zip_cvs_stream(names),
        media_type="application/zip",
        headers={"Content-Disposition": f'attachment; filename="cv-screener-{stamp}.zip"'},
    )


class ScreeningReq(BaseModel):
    profile: str
    session_id: str | None = None
    cvs: list[str] | None = None


class CreateSessionReq(BaseModel):
    name: str
    profile: str
    cvs: list[str]


@app.get("/api/jobs")
def api_list_jobs():
    return {"jobs": screener.get_jobs()}


@app.post("/api/jobs")
def api_start_job(req: ScreeningReq):
    if not req.profile or not req.profile.strip():
        raise HTTPException(status_code=422, detail="Inserisci il profilo target")
    if req.session_id:
        session = screener.get_session(req.session_id)
        if session is None:
            raise HTTPException(status_code=404, detail="Posizione non trovata")
        cvs = req.cvs if req.cvs else session.get("cvs", [])
        if not cvs:
            raise HTTPException(status_code=422, detail="La posizione non ha CV selezionati")
        job = screener.start_job(req.profile.strip(), session_id=req.session_id, cv_names=cvs)
    else:
        if not screener.list_cvs():
            raise HTTPException(status_code=422, detail="Nessun CV presente nella cartella CVs")
        job = screener.start_job(req.profile.strip())
    return {"job": job}


@app.get("/api/jobs/{job_id}")
def api_get_job(job_id: str):
    job = screener.get_job(job_id)
    if job is None:
        raise HTTPException(status_code=404, detail="Job non trovato")
    return {"job": job}


# ── Sessioni (posizioni) ─────────────────────────────
@app.get("/api/sessions")
def api_list_sessions():
    return {"sessions": screener.list_sessions()}


@app.post("/api/sessions")
def api_create_session(req: CreateSessionReq):
    if not req.name or not req.name.strip():
        raise HTTPException(status_code=422, detail="Inserisci il nome della posizione")
    return screener.create_session(req.name, req.profile, req.cvs)


@app.delete("/api/sessions/{session_id}")
def api_delete_session(session_id: str):
    if not screener.delete_session(session_id):
        raise HTTPException(status_code=404, detail="Posizione non trovata")
    return {"ok": True}


# ── Report ────────────────────────────────────────────
@app.get("/api/reports")
def api_list_reports(session: str | None = None):
    return {"reports": screener.list_reports(session)}


@app.get("/api/reports/summary")
def api_reports_summary(session: str | None = None):
    """Dati strutturati dei candidati (voto, giudizio, esperienza) per la tabella interattiva."""
    return {"candidates": screener.get_candidates_summary(session)}


@app.get("/api/reports/summary/csv")
def api_reports_csv(session: str | None = None):
    """Esporta la classifica in CSV."""
    import csv
    import io as _io

    candidates = screener.get_candidates_summary(session)
    buf = _io.StringIO()
    writer = csv.writer(buf, delimiter=";")
    writer.writerow(["Nome", "Voto", "Giudizio", "Esperienza", "Fit tecnico"])
    for c in candidates:
        writer.writerow([c["name"], c.get("score", ""), c.get("verdict", ""), c.get("experience", ""), c.get("fit", "")])
    buf.seek(0)
    return StreamingResponse(
        iter([buf.getvalue()]),
        media_type="text/csv; charset=utf-8",
        headers={"Content-Disposition": 'attachment; filename="classifica.csv"'},
    )


@app.get("/api/reports/{name}")
def api_get_report(name: str, session: str | None = None):
    content = screener.read_report(name, session)
    if content is None:
        raise HTTPException(status_code=404, detail="Report non trovato")
    return {"name": name, "content": content}


@app.get("/api/reports/{name}/pdf")
def api_report_pdf(name: str, session: str | None = None):
    """Converte un report markdown in PDF e lo restituisce come download."""
    content = screener.read_report(name, session)
    if content is None:
        raise HTTPException(status_code=404, detail="Report non trovato")
    body = markdown.markdown(content, extensions=["tables", "fenced_code"])
    html = f"""<!DOCTYPE html>
<html><head><meta charset="utf-8">
<style>
body {{ font-family: Arial, sans-serif; font-size: 11px; padding: 20px; }}
h1 {{ font-size: 18px; }} h2 {{ font-size: 14px; }} h3 {{ font-size: 12px; }}
table {{ border-collapse: collapse; width: 100%; }}
th, td {{ border: 1px solid #ccc; padding: 6px; text-align: left; font-size: 10px; }}
th {{ background: #f5f5f5; }}
code {{ font-family: Consolas, monospace; font-size: 10px; }}
pre {{ background: #f5f5f5; padding: 8px; font-size: 10px; }}
</style></head><body>{body}</body></html>"""
    buf = io.BytesIO()
    pisa.CreatePDF(html, dest=buf, encoding="utf-8")
    buf.seek(0)
    return StreamingResponse(
        buf,
        media_type="application/pdf",
        headers={"Content-Disposition": f'attachment; filename="{name}.pdf"'},
    )


# ── Indeed ────────────────────────────────────────────
INDEED_DIR = Path.home() / "indeedBulkResumesDownloader"
INDEED_DOWNLOADS = INDEED_DIR / "downloads"
INDEED_PYTHON = INDEED_DIR / ".venv" / "Scripts" / "python.exe"


@app.post("/api/indeed/launch")
def api_indeed_launch():
    """Apri una nuova finestra terminale con il downloader Indeed."""
    if not INDEED_DIR.is_dir():
        raise HTTPException(status_code=404, detail="Cartella indeedBulkResumesDownloader non trovata in " + str(Path.home()))
    script = INDEED_DIR / "indeed_downloader.py"
    if not script.is_file():
        raise HTTPException(status_code=404, detail="indeed_downloader.py non trovato")
    python = str(INDEED_PYTHON) if INDEED_PYTHON.is_file() else "python"
    runner = str(Path(__file__).resolve().parent / "indeed_runner.py")
    cmd = f'start cmd /k "cd /d {INDEED_DIR} && {python} {runner} {script}"'
    subprocess.Popen(cmd, shell=True)
    return {"ok": True, "message": "Finestra terminale aperta. Loggati su Indeed Employer e segui il menu."}


@app.post("/api/indeed/sync")
def api_indeed_sync():
    """Copia i PDF scaricati da Indeed nella cartella CVs."""
    if not INDEED_DOWNLOADS.is_dir():
        raise HTTPException(status_code=404, detail="Cartella downloads non trovata. Esegui prima il download da Indeed.")
    copied = []
    for pdf in INDEED_DOWNLOADS.rglob("*.pdf"):
        target = screener.DEFAULT_CV_DIR / pdf.name
        if not target.exists():
            shutil.copy2(pdf, target)
            copied.append(pdf.name)
    return {"copied": copied, "count": len(copied), "cvs": screener.list_cvs()}


def _indeed_pdf_count() -> int:
    if not INDEED_DOWNLOADS.is_dir():
        return 0
    return sum(1 for p in INDEED_DOWNLOADS.rglob("*.pdf") if p.is_file())


@app.get("/api/indeed/downloads")
def api_indeed_downloads():
    """Numero di PDF presenti nell'archivio locale del downloader Indeed."""
    return {"count": _indeed_pdf_count(), "available": INDEED_DOWNLOADS.is_dir()}


@app.delete("/api/indeed/downloads")
def api_indeed_clear_downloads():
    """Elimina definitivamente tutti i PDF scaricati da Indeed (nessun cestino).

    Opera solo sulla cartella fissa INDEED_DOWNLOADS, nessun percorso esterno.
    """
    if not INDEED_DOWNLOADS.is_dir():
        raise HTTPException(status_code=404, detail="Cartella downloads non trovata. Nessun archivio da svuotare.")
    deleted = 0
    for pdf in INDEED_DOWNLOADS.rglob("*.pdf"):
        if not pdf.is_file():
            continue
        try:
            pdf.unlink()
            deleted += 1
        except OSError:
            continue
    return {"deleted": deleted, "count": _indeed_pdf_count()}


# ── Frontend statico (build) ──────────────────────────
if FRONTEND_DIST.is_dir():
    app.mount("/", StaticFiles(directory=str(FRONTEND_DIST), html=True), name="frontend")


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(app, host="127.0.0.1", port=8000)